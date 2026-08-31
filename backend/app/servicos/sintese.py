"""Matriz de sintese e perguntas sobre o conjunto das analises.

A diferenca para `analise_ia` e o insumo: la o contexto e UM PDF, aqui sao as
respostas ja extraidas de VARIOS artigos. Nenhum PDF e reenviado - o que vai
para o modelo e a matriz, que e o proprio produto da revisao.

Isso tambem define o limite honesto da funcionalidade: as lacunas apontadas
valem para o que a matriz cobre, nao para a literatura da area. Se o corpus
tem 12 artigos, a resposta e sobre esses 12.
"""

from __future__ import annotations

import csv
import io
import logging
from datetime import datetime, timezone
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from .. import config
from ..modelos import Artigo, Pergunta, Resposta
from . import analise_claude_code

LOG = logging.getLogger(__name__)

# Acima disso o prompt fica grande demais para uma unica chamada. Com ~8
# perguntas por artigo, 120 artigos ja passam de 100 mil tokens.
MAX_ARTIGOS_NO_CONTEXTO = 120

MARCA_IA = "I.A:"


class SinteseError(Exception):
    """Falha que o usuario precisa ver na tela."""


def montar_matriz(
    sessao: Session,
    busca_id: int | None = None,
    somente_analisados: bool = True,
) -> dict[str, Any]:
    """Monta a matriz artigos x perguntas."""
    perguntas = list(
        sessao.scalars(
            select(Pergunta)
            .where(Pergunta.ativa.is_(True))
            .order_by(Pergunta.ordem, Pergunta.id)
        )
    )

    filtros = []
    if busca_id is not None:
        filtros.append(Artigo.busca_id == busca_id)
    if somente_analisados:
        filtros.append(Artigo.analisado.is_(True))

    artigos = list(
        sessao.scalars(
            select(Artigo).where(*filtros).order_by(Artigo.ano, Artigo.id)
        )
    )
    if not artigos:
        return {"perguntas": [], "artigos": [], "total": 0}

    ids = [a.id for a in artigos]
    respostas: dict[tuple[int, int], str] = {
        (r.artigo_id, r.pergunta_id): r.texto
        for r in sessao.scalars(
            select(Resposta).where(Resposta.artigo_id.in_(ids))
        )
    }

    linhas = [
        {
            "artigo_id": a.id,
            "titulo": a.titulo,
            "autores": a.autores or [],
            "ano": a.ano,
            "venue": a.venue,
            "doi": a.doi,
            "analisado": bool(a.analisado),
            "respostas": {
                p.id: (respostas.get((a.id, p.id)) or "") for p in perguntas
            },
        }
        for a in artigos
    ]

    return {
        "perguntas": [
            {"id": p.id, "texto": p.texto, "ordem": p.ordem} for p in perguntas
        ],
        "artigos": linhas,
        "total": len(linhas),
    }


def matriz_para_csv(matriz: dict[str, Any]) -> str:
    """Exporta a matriz no formato que vai para a planilha da dissertacao."""
    saida = io.StringIO()
    escritor = csv.writer(saida, delimiter=";", quoting=csv.QUOTE_ALL)

    cabecalho = ["Titulo", "Autores", "Ano", "Veiculo", "DOI"]
    cabecalho += [p["texto"] for p in matriz["perguntas"]]
    escritor.writerow(cabecalho)

    for linha in matriz["artigos"]:
        registro = [
            linha["titulo"],
            "; ".join(linha["autores"]),
            linha["ano"] or "",
            linha["venue"] or "",
            linha["doi"] or "",
        ]
        registro += [
            linha["respostas"].get(p["id"], "") for p in matriz["perguntas"]
        ]
        escritor.writerow(registro)

    # BOM para o Excel abrir os acentos certos ao dar duplo clique no arquivo.
    return "﻿" + saida.getvalue()


def _formatar_para_prompt(matriz: dict[str, Any]) -> str:
    """Serializa a matriz de um jeito que o modelo consiga navegar.

    Marca de onde veio cada resposta. O texto que o pesquisador escreveu vale
    mais que a sugestao da maquina, e o modelo precisa saber a diferenca antes
    de tratar tudo como evidencia do mesmo peso.
    """
    partes: list[str] = []
    for i, linha in enumerate(matriz["artigos"], 1):
        cabecalho = f"### Artigo {i}: {linha['titulo']}"
        if linha.get("ano"):
            cabecalho += f" ({linha['ano']})"
        partes.append(cabecalho)
        if linha.get("venue"):
            partes.append(f"Veiculo: {linha['venue']}")

        for p in matriz["perguntas"]:
            texto = (linha["respostas"].get(p["id"]) or "").strip()
            if not texto:
                continue
            origem = (
                "extraida por IA" if texto.lstrip().startswith(MARCA_IA)
                else "escrita pelo pesquisador"
                if MARCA_IA not in texto
                else "escrita pelo pesquisador + complemento de IA"
            )
            partes.append(f"- **{p['texto']}** ({origem})")
            partes.append(f"  {texto}")
        partes.append("")
    return "\n".join(partes)


PREAMBULO = """\
Voce esta ajudando um pesquisador de mestrado a escrever um projeto de \
pesquisa sobre video imersivo e volumetrico, streaming e Head Mounted \
Displays (HMD).

Abaixo esta a matriz de sintese da revisao bibliografica dele: os artigos ja \
analisados e, para cada um, as respostas as perguntas de pesquisa que ele \
definiu.

Regras:

1. Responda com base NESTA matriz. Se algo nao esta nela, diga que a matriz \
nao cobre, em vez de completar com conhecimento geral da area.
2. Ao afirmar algo, cite quais artigos sustentam - pelo titulo ou pelo numero \
que aparece na matriz.
3. Distinga o que a matriz mostra do que voce esta inferindo. Marque inferencia \
como inferencia.
4. Uma lacuna so e lacuna se voce puder apontar a celula vazia ou o padrao \
ausente que a revela. "Ninguem estudou X" sem base na matriz nao serve.
5. Lembre que a matriz e uma amostra, nao a literatura inteira. Uma ausencia \
aqui pode significar que o recorte da busca nao alcancou o tema - diga isso \
quando for plausivel.
6. Escreva em portugues do Brasil, direto, sem preambulo de cortesia.
"""


def perguntar(matriz: dict[str, Any], pergunta: str) -> tuple[str, dict[str, Any]]:
    """Faz uma pergunta sobre o conjunto. Devolve (resposta, metadados)."""
    if not matriz["artigos"]:
        raise SinteseError(
            "Nenhum artigo analisado ainda. Responda as perguntas de pesquisa "
            "de pelo menos um artigo antes de consultar a sintese."
        )
    if len(matriz["artigos"]) > MAX_ARTIGOS_NO_CONTEXTO:
        raise SinteseError(
            f"A matriz tem {len(matriz['artigos'])} artigos e passa do limite de "
            f"{MAX_ARTIGOS_NO_CONTEXTO} por consulta. Filtre por linha de "
            "pesquisa para reduzir o conjunto."
        )

    prompt = "\n".join(
        [
            PREAMBULO,
            "---",
            "## Matriz de sintese",
            "",
            _formatar_para_prompt(matriz),
            "---",
            "## Pergunta do pesquisador",
            "",
            pergunta.strip(),
        ]
    )

    if config.MODO_ANALISE == "claude_code":
        cli = analise_claude_code.localizar_cli(config.CLAUDE_CLI)
        if cli is None:
            raise SinteseError(analise_claude_code.INSTRUCAO_INSTALACAO)
        try:
            texto, _estruturado, metadados = analise_claude_code.rodar(
                prompt=prompt,
                cli=cli,
                timeout_s=config.TIMEOUT_ANALISE_S,
                modelo=config.MODELO_ANALISE_CLI,
                # Tudo o que a resposta precisa ja esta no prompt; sem
                # ferramenta nenhuma a sessao fica mais barata e mais previsivel.
                ferramentas=(),
            )
        except RuntimeError as exc:
            raise SinteseError(str(exc)) from exc
        metadados["modelo"] = "Claude Code (assinatura)"
        return texto.strip(), metadados

    return _perguntar_via_api(prompt)


def _perguntar_via_api(prompt: str) -> tuple[str, dict[str, Any]]:
    import anthropic

    try:
        cliente = anthropic.Anthropic()
        with cliente.messages.stream(
            model=config.MODELO_ANALISE,
            max_tokens=16000,
            messages=[{"role": "user", "content": prompt}],
            thinking={"type": "adaptive"},
        ) as fluxo:
            mensagem = fluxo.get_final_message()
    except TypeError as exc:
        if "authentication" in str(exc).lower():
            raise SinteseError(
                "Nenhuma credencial da Anthropic encontrada. Defina "
                "ANTHROPIC_API_KEY no .env, ou use MODO_ANALISE=claude_code."
            ) from exc
        raise
    except anthropic.APIError as exc:
        raise SinteseError(f"A API da Anthropic falhou: {exc}") from exc

    if mensagem.stop_reason == "refusal":
        raise SinteseError("O modelo recusou responder a esta consulta.")

    texto = "".join(b.text for b in mensagem.content if b.type == "text")
    entrada = mensagem.usage.input_tokens
    saida = mensagem.usage.output_tokens
    return texto.strip(), {
        "modelo": config.MODELO_ANALISE,
        "tokens_entrada": entrada,
        "tokens_saida": saida,
        "custo_usd": entrada / 1_000_000 * 5.0 + saida / 1_000_000 * 25.0,
    }


# --- Exportacao para planilha ---------------------------------------------

# Largura das colunas em "caracteres" do Excel. As respostas sao textos de
# varias frases, entao precisam de bem mais espaco que os metadados.
LARGURA_TITULO = 46
LARGURA_META = 16
LARGURA_RESPOSTA = 52


def matriz_para_xlsx(matriz: dict[str, Any], escopo: str) -> bytes:
    """Gera o .xlsx da matriz.

    Duas abas: `Matriz` com os dados e `Resumo` com a procedencia (quando foi
    exportado, que filtro estava ligado, quantos artigos). Numa dissertacao a
    planilha circula solta do sistema, e sem esse registro ninguem consegue
    dizer de qual recorte ela saiu.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    livro = Workbook()
    aba = livro.active
    aba.title = "Matriz"

    colunas_meta = ["Título", "Autores", "Ano", "Veículo", "DOI"]
    cabecalho = colunas_meta + [p["texto"] for p in matriz["perguntas"]]
    aba.append(cabecalho)

    fundo = PatternFill("solid", fgColor="1F2937")
    for indice, _ in enumerate(cabecalho, start=1):
        celula = aba.cell(row=1, column=indice)
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = fundo
        celula.alignment = Alignment(vertical="top", wrap_text=True)

    for linha in matriz["artigos"]:
        aba.append(
            [
                linha["titulo"],
                "; ".join(linha["autores"]),
                linha["ano"] or "",
                linha["venue"] or "",
                linha["doi"] or "",
                *[linha["respostas"].get(p["id"], "") for p in matriz["perguntas"]],
            ]
        )

    # Texto quebrado e alinhado no topo: sem isso uma resposta de mil
    # caracteres vira uma linha unica cortada na borda da coluna.
    alinhamento = Alignment(vertical="top", wrap_text=True)
    for fileira in aba.iter_rows(min_row=2):
        for celula in fileira:
            celula.alignment = alinhamento

    for indice in range(1, len(cabecalho) + 1):
        letra = get_column_letter(indice)
        if indice == 1:
            aba.column_dimensions[letra].width = LARGURA_TITULO
        elif indice <= len(colunas_meta):
            aba.column_dimensions[letra].width = LARGURA_META
        else:
            aba.column_dimensions[letra].width = LARGURA_RESPOSTA

    # Congela o cabecalho e a coluna do titulo: rolando ate a 8a pergunta,
    # da para continuar vendo de qual artigo e a linha.
    aba.freeze_panes = "B2"
    aba.auto_filter.ref = (
        f"A1:{get_column_letter(len(cabecalho))}{aba.max_row}"
    )

    resumo = livro.create_sheet("Resumo")
    resumo.column_dimensions["A"].width = 24
    resumo.column_dimensions["B"].width = 96
    linhas_resumo = [
        ("Exportado em", datetime.now(timezone.utc).astimezone().strftime("%d/%m/%Y %H:%M")),
        ("Escopo", escopo),
        ("Artigos", len(matriz["artigos"])),
        ("Perguntas", len(matriz["perguntas"])),
        ("", ""),
        ("Perguntas de pesquisa", ""),
    ]
    for rotulo, valor in linhas_resumo:
        resumo.append([rotulo, valor])
    for indice, pergunta in enumerate(matriz["perguntas"], start=1):
        resumo.append([f"{indice}.", pergunta["texto"]])

    for fileira in resumo.iter_rows():
        for celula in fileira:
            celula.alignment = Alignment(vertical="top", wrap_text=True)
    for indice in (1, 6):
        resumo.cell(row=indice, column=1).font = Font(bold=True)

    fluxo = io.BytesIO()
    livro.save(fluxo)
    return fluxo.getvalue()
